"""Visible-table XLSX codec for the authoritative SEER ledgers.

The workbook is deliberately a normal, editable spreadsheet.  Every business
value is in a visible Excel table; there is no hidden JSON worksheet or custom
property which is required to reconstruct the ledger.  A small, visible type
column accompanies each field so that ``null``, empty strings, booleans and
integer money values remain distinguishable after an Excel round trip.

This module is pure local encoding/decoding.  It performs no SharePoint
operation and is consequently safe to use from migration tooling and tests.
"""

from __future__ import annotations

import hashlib
import json
import math
import posixpath
from collections.abc import Mapping
from datetime import date, datetime, time
from io import BytesIO
from copy import copy
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


EXPENSE_WORKBOOK_KIND = "expense"
FINANCE_WORKBOOK_KIND = "finance"
WORKBOOK_SCHEMA_VERSION = 1

_EXPENSE_SHEETS = ("Expenses", "Events", "Collisions", "Evidence")
_FINANCE_SHEETS = ("Transactions",)
_METADATA_SHEET = "Workbook"
_OPTIONAL_EXPENSE_SHEETS = frozenset({"Receipt items"})
_TYPE_SUFFIX = " [type]"
_TYPE_VALUES = frozenset({
    "missing", "null", "str", "int", "float", "bool", "date", "datetime", "time", "json",
})
_MISSING = object()


class WorkbookCodecError(ValueError):
    """The XLSX is not a valid visible-table SEER ledger workbook."""


class WorkbookValidationError(WorkbookCodecError):
    """An editable workbook contains a malformed value or table edit."""


def _kind_sheets(kind: str) -> tuple[str, ...]:
    if kind == EXPENSE_WORKBOOK_KIND:
        return _EXPENSE_SHEETS
    if kind == FINANCE_WORKBOOK_KIND:
        return _FINANCE_SHEETS
    raise WorkbookCodecError(f"unknown workbook kind {kind!r}")


def _optional_sheets(kind: str) -> frozenset[str]:
    """Return approved visible auxiliary sheets preserved outside ledger state."""
    if kind == EXPENSE_WORKBOOK_KIND:
        return _OPTIONAL_EXPENSE_SHEETS
    if kind == FINANCE_WORKBOOK_KIND:
        return frozenset()
    raise WorkbookCodecError(f"unknown workbook kind {kind!r}")


def _json_value(value: Any) -> Any:
    """Return a deterministic JSON-compatible representation for hashing."""
    if isinstance(value, Mapping):
        return {str(key): _json_value(item) for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        raise WorkbookCodecError("non-finite floating point values are not supported")
    return value


def semantic_payload_hash(payload: Mapping[str, Any]) -> str:
    """Hash ledger meaning, excluding ZIP/XML metadata Excel may rewrite."""
    encoded = json.dumps(
        _json_value(payload),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def workbook_semantic_hash(content: bytes, *, kind: str | None = None) -> str:
    """Decode an XLSX and hash its visible table meaning."""
    return WorkbookCodec.semantic_hash(content, kind=kind)


def _type_of(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, str):
        return "str"
    if isinstance(value, (dict, list, tuple)):
        return "json"
    raise WorkbookCodecError(
        f"unsupported value type {type(value).__name__}; only scalar values are "
        "editable in workbook tables"
    )


def _cell_value(value: Any, kind: str) -> Any:
    if kind == "null":
        return None
    if kind == "str" and isinstance(value, str) and value.startswith("="):
        raise WorkbookCodecError(
            "formula strings are not supported; enter a literal text value"
        )
    if kind == "json":
        return json.dumps(_json_value(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return value


def _decode_cell(value: Any, declared: Any, *, sheet: str, row: int, field: str) -> Any:
    if not isinstance(declared, str) or declared not in _TYPE_VALUES:
        raise WorkbookValidationError(
            f"{sheet}!{field}{row}: type marker must be one of {sorted(_TYPE_VALUES)}"
        )
    if isinstance(value, str) and value.startswith("="):
        raise WorkbookValidationError(
            f"{sheet}!{field}{row}: formulas are not supported; enter a calculated value"
        )
    if declared == "missing":
        if value not in (None, ""):
            # A normal spreadsheet edit does not require the operator to
            # understand or update the adjacent marker column.  Infer a safe
            # scalar type when a formerly missing cell receives a value.
            return _ordinary_value(value, sheet=sheet, row=row, field=field)
        return _MISSING
    if declared == "null":
        if value not in (None, ""):
            return _ordinary_value(value, sheet=sheet, row=row, field=field)
        return None
    if declared == "str":
        # openpyxl returns None for an explicitly saved empty string.  The
        # marker retains that distinction from a null cell.
        if value is None:
            return ""
        if not isinstance(value, str):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected text")
        return value
    if declared == "bool":
        if not isinstance(value, bool):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected boolean")
        return value
    if declared == "int":
        if isinstance(value, bool) or not isinstance(value, int):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected integer")
        return value
    if declared == "float":
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected number")
        return float(value)
    if declared in {"date", "datetime", "time"}:
        expected = {"date": date, "datetime": datetime, "time": time}[declared]
        if not isinstance(value, expected):
            # A user may edit an ISO value into a text cell.  Reject it rather
            # than silently changing the source type.
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected Excel {declared}")
        return value
    if declared == "json":
        if not isinstance(value, str):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: expected JSON text")
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError as exc:
            raise WorkbookValidationError(f"{sheet}!{field}{row}: malformed JSON text: {exc.msg}") from exc
        if not isinstance(parsed, (dict, list)):
            raise WorkbookValidationError(f"{sheet}!{field}{row}: JSON value must be an object or array")
        return parsed
    raise AssertionError(declared)


def _ordinary_value(value: Any, *, sheet: str, row: int, field: str) -> Any:
    """Validate a value entered into a blank/null cell without a marker edit."""
    try:
        kind = _type_of(value)
    except WorkbookCodecError as exc:
        raise WorkbookValidationError(f"{sheet}!{field}{row}: {exc}") from exc
    if kind == "str" and isinstance(value, str) and value.startswith("="):
        raise WorkbookValidationError(
            f"{sheet}!{field}{row}: formulas are not supported; enter a calculated value"
        )
    return value


def _iso_value(value: Any, kind: str) -> Any:
    if kind in {"date", "datetime", "time"}:
        return value.isoformat()
    return value


def _table_name(sheet: str) -> str:
    return "T" + "".join(character if character.isalnum() else "_" for character in sheet)


def _set_widths(sheet: Any, headers: list[str]) -> None:
    for index, header in enumerate(headers, 1):
        width = min(max(len(header) + 2, 12), 32)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _add_table(sheet: Any, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    # A one-row table is valid for an empty source list and remains editable.
    ref = f"A1:{get_column_letter(max(1, len(headers)))}{max(1, len(rows) + 1)}"
    table = Table(displayName=_table_name(sheet.title), ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    _set_widths(sheet, headers)


def _replace_table(
    sheet: Any,
    headers: list[str],
    rows: list[list[Any]],
    *,
    table_name: str,
) -> None:
    """Replace table values while retaining the existing table/style object."""
    tables = list(sheet.tables.values())
    if len(tables) != 1:
        raise WorkbookValidationError(f"{table_name}: exactly one visible Excel table is required")
    table = tables[0]
    source_rows = list(sheet.iter_rows(values_only=False))
    source_header = [cell.value for cell in source_rows[0]] if source_rows else []
    while source_header and source_header[-1] is None:
        source_header.pop()
    if source_header != headers:
        # Empty placeholder tables are allowed to acquire their first real
        # fields.  Other column changes could discard human formatting and are
        # deliberately rejected by WorkbookCodec.update.
        if source_header != ["_empty", "_empty" + _TYPE_SUFFIX]:
            raise WorkbookValidationError(
                f"{table_name}: source table columns changed unexpectedly"
            )
        for index, header in enumerate(headers, 1):
            sheet.cell(1, index).value = header
    old_row_count = max(0, sheet.max_row - 1)
    new_row_count = len(rows)
    if old_row_count > new_row_count:
        sheet.delete_rows(new_row_count + 2, old_row_count - new_row_count)
    # Copy the final existing data row's formatting into newly created rows.
    template_row = min(max(2, sheet.max_row), max(2, new_row_count + 1))
    template_cells = [
        copy(sheet.cell(template_row, index))
        for index in range(1, len(headers) + 1)
    ] if sheet.max_row >= 2 else []
    for row_index, values in enumerate(rows, 2):
        for column_index, value in enumerate(values, 1):
            target = sheet.cell(row_index, column_index)
            if row_index > old_row_count + 1 and column_index <= len(template_cells):
                template = template_cells[column_index - 1]
                target._style = copy(template._style)
                target.number_format = template.number_format
                target.alignment = copy(template.alignment)
                target.protection = copy(template.protection)
                target.fill = copy(template.fill)
                target.border = copy(template.border)
                target.font = copy(template.font)
            target.value = value
    # A table with no records is represented by its header row only.
    end_row = max(1, new_row_count + 1)
    table.ref = f"A1:{get_column_letter(len(headers))}{end_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


def _headers_preserve_source_columns(
    source: list[str], desired: list[str]
) -> bool:
    """Allow source-only columns whose decoded rows are all ``missing``."""
    if len(source) % 2 or len(desired) % 2:
        return False
    source_fields = source[::2]
    desired_fields = desired[::2]
    source_types = source[1::2]
    desired_types = desired[1::2]
    if any(
        source_types[index] != field + _TYPE_SUFFIX
        for index, field in enumerate(source_fields)
    ):
        return False
    if any(
        desired_types[index] != field + _TYPE_SUFFIX
        for index, field in enumerate(desired_fields)
    ):
        return False
    position = 0
    for field in desired_fields:
        try:
            position = source_fields.index(field, position) + 1
        except ValueError:
            return False
    return True


def _stable_xlsx_bytes(content: bytes) -> bytes:
    """Normalize ZIP member order/timestamps without changing XML payloads."""
    source = BytesIO(content)
    output = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(
        output, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as normalized:
        for name in sorted(archive.namelist()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            original = archive.getinfo(name)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            normalized.writestr(info, archive.read(name))
    return output.getvalue()


_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_WORKBOOK_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _package_target(source_part: str, target: str) -> str:
    """Resolve an OOXML relationship target to a ZIP member name."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _relationship_part(source_part: str) -> str:
    directory, filename = posixpath.split(source_part)
    return posixpath.join(directory, "_rels", filename + ".rels")


def _editable_package_parts(content: bytes, kind: str) -> set[str]:
    """Return only worksheet/table members changed by a visible-table update.

    Relationship files and package content types are deliberately not in this
    set.  The source package remains authoritative for those files, including
    relationships to package features that openpyxl does not understand.
    """
    try:
        with ZipFile(BytesIO(content), "r") as archive:
            names = set(archive.namelist())
            workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            workbook_rels = ElementTree.fromstring(
                archive.read("xl/_rels/workbook.xml.rels")
            )
            relationship_targets = {
                relationship.attrib["Id"]: _package_target(
                    "xl/workbook.xml", relationship.attrib["Target"]
                )
                for relationship in workbook_rels.findall(
                    f"{{{_PACKAGE_REL_NS}}}Relationship"
                )
                if relationship.attrib.get("Type") == f"{_DOCUMENT_REL_NS}/officeDocument"
                or relationship.attrib.get("Type") == f"{_DOCUMENT_REL_NS}/worksheet"
            }
            requested_sheets = set(_kind_sheets(kind)) | {_METADATA_SHEET}
            parts: set[str] = set()
            for sheet in workbook_xml.findall(
                f"{{{_WORKBOOK_NS}}}sheets/{{{_WORKBOOK_NS}}}sheet"
            ):
                if sheet.attrib.get("name") not in requested_sheets:
                    continue
                relationship_id = sheet.attrib.get(f"{{{_DOCUMENT_REL_NS}}}id")
                worksheet_part = relationship_targets.get(relationship_id or "")
                if not worksheet_part or worksheet_part not in names:
                    raise WorkbookCodecError(
                        f"workbook sheet {sheet.attrib.get('name')!r} has no valid worksheet part"
                    )
                parts.add(worksheet_part)
                worksheet_rels_name = _relationship_part(worksheet_part)
                if worksheet_rels_name not in names:
                    raise WorkbookCodecError(
                        f"worksheet {worksheet_part!r} has no relationship part"
                    )
                worksheet_rels = ElementTree.fromstring(
                    archive.read(worksheet_rels_name)
                )
                for relationship in worksheet_rels.findall(
                    f"{{{_PACKAGE_REL_NS}}}Relationship"
                ):
                    if relationship.attrib.get("Type") == f"{_DOCUMENT_REL_NS}/table":
                        table_part = _package_target(
                            worksheet_part, relationship.attrib["Target"]
                        )
                        if table_part not in names:
                            raise WorkbookCodecError(
                                f"worksheet {worksheet_part!r} references missing table {table_part!r}"
                            )
                        parts.add(table_part)
            return parts
    except (BadZipFile, OSError) as exc:
        raise WorkbookCodecError(f"invalid XLSX package: {exc}") from exc


def _merge_updated_package(
    source: bytes, updated: bytes, *, editable_parts: set[str]
) -> bytes:
    """Use openpyxl output only for edited worksheets/tables.

    The source archive is the authority for every other member.  This is
    stronger than checking for dropped members after an openpyxl save: an
    unknown relationship, content-type declaration, custom XML part, or
    embedded package is copied byte-for-byte rather than merely allowed to
    survive if openpyxl happens not to rewrite it.
    """
    try:
        with ZipFile(BytesIO(source), "r") as original, ZipFile(
            BytesIO(updated), "r"
        ) as rewritten:
            source_names = set(original.namelist())
            updated_names = set(rewritten.namelist())
            missing_editable = sorted(editable_parts - updated_names)
            if missing_editable:
                raise WorkbookCodecError(
                    "updated XLSX package is missing edited parts: "
                    + ", ".join(missing_editable)
                )
            unexpected = sorted(updated_names - source_names)
            if unexpected:
                raise WorkbookCodecError(
                    "automated update introduced unexpected XLSX package parts: "
                    + ", ".join(unexpected)
                )
            output = BytesIO()
            with ZipFile(
                output, "w", compression=ZIP_DEFLATED, compresslevel=9
            ) as merged:
                for name in sorted(source_names):
                    original_info = original.getinfo(name)
                    info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                    info.compress_type = ZIP_DEFLATED
                    info.external_attr = original_info.external_attr
                    info.create_system = original_info.create_system
                    payload = (
                        rewritten.read(name)
                        if name in editable_parts
                        else original.read(name)
                    )
                    merged.writestr(info, payload)
            return output.getvalue()
    except (BadZipFile, OSError) as exc:
        raise WorkbookCodecError(f"updated XLSX package is invalid: {exc}") from exc


def _field_order(rows: list[Mapping[str, Any]]) -> list[str]:
    fields: list[str] = []
    known: set[str] = set()
    for row in rows:
        for field in row:
            if not isinstance(field, str) or not field:
                raise WorkbookCodecError("table field names must be non-empty strings")
            if field not in known:
                known.add(field)
                fields.append(field)
    return fields


def _table_headers(rows: list[Mapping[str, Any]]) -> list[str]:
    fields = _field_order(rows)
    headers: list[str] = []
    for field in fields:
        if field.endswith(_TYPE_SUFFIX):
            raise WorkbookCodecError(f"field name {field!r} is reserved for codec type markers")
        headers.extend((field, field + _TYPE_SUFFIX))
    # Empty lists have no source fields; a visible placeholder is still a real
    # table and is removed by the decoder.
    return headers or ["_empty", "_empty" + _TYPE_SUFFIX]


def _table_rows(rows: list[Mapping[str, Any]], headers: list[str]) -> list[list[Any]]:
    fields = [header for header in headers if not header.endswith(_TYPE_SUFFIX) and header != "_empty"]
    result: list[list[Any]] = []
    for source in rows:
        row: list[Any] = []
        for field in fields:
            if field not in source:
                value = None
                kind = "missing"
            else:
                value = source[field]
                kind = _type_of(value)
            row.extend((_cell_value(value, kind), kind))
        result.append(row)
    return result


def _metadata_rows(payload: Mapping[str, Any], kind: str) -> list[list[Any]]:
    rows: list[list[Any]] = [
        ["document_kind", kind, "str"],
        ["codec_schema_version", WORKBOOK_SCHEMA_VERSION, "int"],
    ]
    for key, value in payload.items():
        if isinstance(value, list):
            continue
        if not isinstance(key, str) or not key:
            raise WorkbookCodecError("top-level metadata keys must be non-empty strings")
        declared = _type_of(value)
        rows.append([key, _cell_value(value, declared), declared])
    return rows


class WorkbookCodec:
    """Encode and decode expense/finance ledgers as visible-table workbooks."""

    @staticmethod
    def encode(payload: Mapping[str, Any], *, kind: str) -> bytes:
        if not isinstance(payload, Mapping):
            raise WorkbookCodecError("workbook payload must be an object")
        sheets = _kind_sheets(kind)
        supported_keys = {
            "expenses": "Expenses",
            "events": "Events",
            "collisions": "Collisions",
            "evidence": "Evidence",
            "transactions": "Transactions",
        }
        unsupported_lists = [
            key for key, value in payload.items()
            if isinstance(value, list) and key not in {
                source_key for source_key, sheet_name in supported_keys.items()
                if sheet_name in sheets
            }
        ]
        if unsupported_lists:
            raise WorkbookCodecError(
                "top-level list field(s) have no visible table: "
                + ", ".join(sorted(str(key) for key in unsupported_lists))
            )
        workbook = Workbook()
        workbook.remove(workbook.active)
        # openpyxl otherwise puts the current clock into core properties.  A
        # deterministic package gives retries the same exact bytes and makes
        # exact-byte journaling useful rather than merely cosmetic.
        fixed_time = datetime(2000, 1, 1)
        workbook.properties.created = fixed_time
        workbook.properties.modified = fixed_time

        metadata = workbook.create_sheet(_METADATA_SHEET)
        _add_table(metadata, ["key", "value", "type"], _metadata_rows(payload, kind))

        for sheet_name in sheets:
            source_key = {
                "Expenses": "expenses",
                "Events": "events",
                "Collisions": "collisions",
                "Evidence": "evidence",
                "Transactions": "transactions",
            }[sheet_name]
            rows = payload.get(source_key, [])
            if not isinstance(rows, list) or not all(isinstance(row, Mapping) for row in rows):
                raise WorkbookCodecError(f"{source_key} must be a list of objects")
            sheet = workbook.create_sheet(sheet_name)
            headers = _table_headers(rows)
            _add_table(sheet, headers, _table_rows(rows, headers))

        for sheet in workbook.worksheets:
            sheet.sheet_view.showGridLines = True
            sheet.sheet_state = "visible"

        output = BytesIO()
        workbook.save(output)
        return _stable_xlsx_bytes(output.getvalue())

    @staticmethod
    def update(content: bytes, payload: Mapping[str, Any], *, kind: str) -> bytes:
        """Apply a ledger state to an existing workbook package.

        This intentionally edits the source workbook rather than rebuilding it.
        Table styles, widths, comments, validations, and other openpyxl-known
        worksheet features therefore survive an automated update.  Package
        features which openpyxl cannot round-trip safely fail closed instead of
        being silently discarded.
        """
        if not isinstance(content, (bytes, bytearray)) or not content:
            raise WorkbookCodecError("source workbook content must be non-empty XLSX bytes")
        WorkbookCodec.decode(bytes(content), kind=kind)
        editable_parts = _editable_package_parts(bytes(content), kind)
        if not isinstance(payload, Mapping):
            raise WorkbookCodecError("workbook payload must be an object")
        workbook = load_workbook(BytesIO(bytes(content)), data_only=False, read_only=False)
        try:
            sheets = _kind_sheets(kind)
            metadata = workbook[_METADATA_SHEET]
            _replace_table(
                metadata,
                ["key", "value", "type"],
                _metadata_rows(payload, kind),
                table_name="Workbook",
            )
            source_keys = {
                "Expenses": "expenses",
                "Events": "events",
                "Collisions": "collisions",
                "Evidence": "evidence",
                "Transactions": "transactions",
            }
            for sheet_name in sheets:
                rows = payload.get(source_keys[sheet_name], [])
                if not isinstance(rows, list) or not all(isinstance(row, Mapping) for row in rows):
                    raise WorkbookCodecError(f"{source_keys[sheet_name]} must be a list of objects")
                sheet = workbook[sheet_name]
                current_headers, _ = _read_table_rows(sheet, sheet_name)
                desired_headers = _table_headers(rows)
                # Keep the source's established columns when a collection is
                # empty; the placeholder has no source meaning and should not
                # erase human formatting/column widths.
                if not rows and current_headers != ["_empty", "_empty" + _TYPE_SUFFIX]:
                    desired_headers = current_headers
                elif current_headers == ["_empty", "_empty" + _TYPE_SUFFIX] and rows:
                    desired_headers = _table_headers(rows)
                elif current_headers != desired_headers:
                    if _headers_preserve_source_columns(current_headers, desired_headers):
                        desired_headers = current_headers
                    else:
                        raise WorkbookValidationError(
                            f"{sheet_name}: structural column edits are unsupported during automated update"
                        )
                _replace_table(
                    sheet,
                    desired_headers,
                    _table_rows(rows, desired_headers),
                    table_name=sheet_name,
                )
            output = BytesIO()
            workbook.save(output)
            updated = output.getvalue()
            merged = _merge_updated_package(
                bytes(content), updated, editable_parts=editable_parts
            )
            return _stable_xlsx_bytes(merged)
        finally:
            workbook.close()

    @staticmethod
    def encode_expense(payload: Mapping[str, Any]) -> bytes:
        return WorkbookCodec.encode(payload, kind=EXPENSE_WORKBOOK_KIND)

    @staticmethod
    def encode_finance(payload: Mapping[str, Any]) -> bytes:
        return WorkbookCodec.encode(payload, kind=FINANCE_WORKBOOK_KIND)

    @staticmethod
    def update_expense(content: bytes, payload: Mapping[str, Any]) -> bytes:
        return WorkbookCodec.update(content, payload, kind=EXPENSE_WORKBOOK_KIND)

    @staticmethod
    def update_finance(content: bytes, payload: Mapping[str, Any]) -> bytes:
        return WorkbookCodec.update(content, payload, kind=FINANCE_WORKBOOK_KIND)

    @staticmethod
    def decode(content: bytes, *, kind: str | None = None) -> dict[str, Any]:
        if not isinstance(content, (bytes, bytearray)) or not content:
            raise WorkbookCodecError("workbook content must be non-empty XLSX bytes")
        try:
            workbook = load_workbook(BytesIO(bytes(content)), data_only=False, read_only=False)
        except (BadZipFile, OSError, ValueError) as exc:
            raise WorkbookCodecError(f"invalid XLSX workbook: {exc}") from exc
        try:
            if any(sheet.sheet_state != "visible" for sheet in workbook.worksheets):
                raise WorkbookValidationError("hidden worksheets are not allowed; all source tables must be visible")
            if _METADATA_SHEET not in workbook.sheetnames:
                raise WorkbookValidationError("missing visible Workbook metadata table")
            metadata = _read_metadata(workbook[_METADATA_SHEET])
            workbook_kind = metadata.get("document_kind")
            if workbook_kind not in {EXPENSE_WORKBOOK_KIND, FINANCE_WORKBOOK_KIND}:
                raise WorkbookValidationError("Workbook.document_kind must be 'expense' or 'finance'")
            if metadata.get("codec_schema_version") != WORKBOOK_SCHEMA_VERSION:
                raise WorkbookValidationError("unsupported workbook codec schema version")
            if kind is not None and workbook_kind != kind:
                raise WorkbookValidationError(
                    f"workbook kind is {workbook_kind!r}, expected {kind!r}"
                )
            required_sheets = set(_kind_sheets(workbook_kind)) | {_METADATA_SHEET}
            allowed_sheets = required_sheets | set(_optional_sheets(workbook_kind))
            actual_sheets = set(workbook.sheetnames)
            if not required_sheets.issubset(actual_sheets) or not actual_sheets.issubset(allowed_sheets):
                unknown = sorted(actual_sheets - allowed_sheets)
                missing = sorted(required_sheets - actual_sheets)
                detail = []
                if missing:
                    detail.append("missing " + ", ".join(missing))
                if unknown:
                    detail.append("unknown " + ", ".join(unknown))
                raise WorkbookValidationError("workbook sheets do not match contract: " + "; ".join(detail))

            payload: dict[str, Any] = {}
            for key, value in metadata.items():
                if key not in {"document_kind", "codec_schema_version"}:
                    payload[key] = value
            payload["schema_version"] = payload.get("schema_version", 1)
            for sheet_name in _kind_sheets(workbook_kind):
                source_key = {
                    "Expenses": "expenses",
                    "Events": "events",
                    "Collisions": "collisions",
                    "Evidence": "evidence",
                    "Transactions": "transactions",
                }[sheet_name]
                payload[source_key] = _read_table(workbook[sheet_name])
            # Keep stable source ordering while placing the schema marker first,
            # matching the structured Markdown representation used by migration.
            ordered: dict[str, Any] = {"schema_version": payload.pop("schema_version")}
            ordered.update(payload)
            return ordered
        finally:
            workbook.close()

    @staticmethod
    def decode_expense(content: bytes) -> dict[str, Any]:
        return WorkbookCodec.decode(content, kind=EXPENSE_WORKBOOK_KIND)

    @staticmethod
    def decode_finance(content: bytes) -> dict[str, Any]:
        return WorkbookCodec.decode(content, kind=FINANCE_WORKBOOK_KIND)

    @staticmethod
    def semantic_hash(content: bytes, *, kind: str | None = None) -> str:
        # Keep this hash independent of optional transport imports.  The queue
        # processor is intentionally deployable without seer-finance, so
        # importing its sibling module here made producer and processor proofs
        # differ depending on sys.path/bootstrap order.
        # Semantic proofs must never bless a workbook which the authoritative
        # decoder would reject (notably formula cells whose cached value is
        # unavailable to this process).
        WorkbookCodec.decode(content, kind=kind)
        from .workbook_semantic import semantic_sha256

        return semantic_sha256(content)


def _read_metadata(sheet: Any) -> dict[str, Any]:
    rows = _read_table_rows(sheet, "Workbook")
    required = {"key", "value", "type"}
    headers, values = rows
    if set(headers) != required:
        raise WorkbookValidationError("Workbook table must have exactly key, value, type columns")
    result: dict[str, Any] = {}
    for index, row in enumerate(values, 2):
        key, value, declared = row
        if isinstance(key, str) and key.startswith("="):
            raise WorkbookValidationError(
                f"Workbook!key{index}: formulas are not supported; enter text"
            )
        if not isinstance(key, str) or not key or key in result:
            raise WorkbookValidationError(f"Workbook!key{index}: duplicate or invalid metadata key")
        result[key] = _decode_cell(value, declared, sheet="Workbook", row=index, field="value")
    return result


def _read_table(sheet: Any) -> list[dict[str, Any]]:
    headers, values = _read_table_rows(sheet, sheet.title)
    if len(headers) % 2:
        raise WorkbookValidationError(f"{sheet.title}: every visible field needs a value and type column")
    fields = headers[::2]
    for index, field in enumerate(fields):
        expected = field + _TYPE_SUFFIX
        if headers[index * 2 + 1] != expected:
            raise WorkbookValidationError(
                f"{sheet.title}: type column for {field!r} must be named {expected!r}"
            )
    if fields == ["_empty"]:
        if values:
            raise WorkbookValidationError(f"{sheet.title}: empty table has unexpected rows")
        return []
    result: list[dict[str, Any]] = []
    for row_number, row in enumerate(values, 2):
        if all(value in (None, "") for value in row):
            raise WorkbookValidationError(f"{sheet.title}!{row_number}: blank rows are not allowed")
        item: dict[str, Any] = {}
        for offset, field in enumerate(fields):
            decoded = _decode_cell(
                row[offset * 2],
                row[offset * 2 + 1],
                sheet=sheet.title,
                row=row_number,
                field=field,
            )
            if decoded is not _MISSING:
                item[field] = decoded
        result.append(item)
    return result


def _read_table_rows(sheet: Any, name: str) -> tuple[list[str], list[list[Any]]]:
    tables = list(sheet.tables.values())
    if len(tables) != 1:
        raise WorkbookValidationError(f"{name}: exactly one visible Excel table is required")
    table = tables[0]
    if table.ref is None or ":" not in table.ref:
        raise WorkbookValidationError(f"{name}: table range is malformed")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise WorkbookValidationError(f"{name}: table has no header row")
    headers = list(rows[0])
    while headers and headers[-1] is None:
        headers.pop()
    if not headers or any(not isinstance(header, str) or not header for header in headers):
        raise WorkbookValidationError(f"{name}: table headers must be non-empty text")
    # Ref must cover the complete used table and must begin at A1.  This makes
    # accidental edits outside the real table fail closed rather than being
    # silently ignored.
    start, end = table.ref.split(":", 1)
    if start != "A1":
        raise WorkbookValidationError(f"{name}: table must start at A1")
    expected_end = f"{get_column_letter(len(headers))}{max(1, len(rows))}"
    if end != expected_end:
        raise WorkbookValidationError(
            f"{name}: table range {table.ref!r} does not cover visible rows (expected {expected_end!r})"
        )
    values: list[list[Any]] = []
    for row in rows[1:]:
        values.append(list(row[: len(headers)]))
        if len(row) > len(headers) and any(value is not None for value in row[len(headers) :]):
            raise WorkbookValidationError(f"{name}: values exist outside the declared table columns")
    return headers, values


# Friendly functional aliases keep callers independent of codec class naming.
encode_expense_workbook = WorkbookCodec.encode_expense
decode_expense_workbook = WorkbookCodec.decode_expense
encode_finance_workbook = WorkbookCodec.encode_finance
decode_finance_workbook = WorkbookCodec.decode_finance
update_expense_workbook = WorkbookCodec.update_expense
update_finance_workbook = WorkbookCodec.update_finance
