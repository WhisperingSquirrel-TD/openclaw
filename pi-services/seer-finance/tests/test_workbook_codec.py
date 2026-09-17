from __future__ import annotations

from io import BytesIO
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

from seer_finance.ledger.workbook_codec import (
    WorkbookCodec,
    WorkbookValidationError,
    _editable_package_parts,
    encode_expense_workbook,
    encode_finance_workbook,
)


class WorkbookCodecTests(unittest.TestCase):
    def test_expense_visible_tables_round_trip_all_source_collections(self) -> None:
        payload = {
            "schema_version": 1,
            "expenses": [{
                "expense_id": "e-1",
                "source_ref": "receipt:1",
                "amount_pence": 1499,
                "empty_string": "",
                "missing_value": None,
                "flag": False,
            }],
            "events": [{
                "event_id": "event-1",
                "expense_id": "e-1",
                "duration_ms": 0,
            }],
            "collisions": [{
                "collision_id": "collision-1",
                "original_facts_json": "{\"amount_pence\":1}",
                "incoming_facts_json": "{\"amount_pence\":2}",
            }],
            "evidence": [{
                "evidence_id": "evidence-1",
                "source_ref": "receipt:1",
                "sha256": None,
            }],
        }
        encoded = encode_expense_workbook(payload)
        decoded = WorkbookCodec.decode_expense(encoded)
        self.assertEqual(payload, decoded)

    def test_finance_integer_money_and_nulls_do_not_become_float_or_text(self) -> None:
        payload = {
            "schema_version": 1,
            "transactions": [{
                "txn_id": "txn-1",
                "date": "2026-08-10",
                "direction": "expense",
                "amount_pence": 1200,
                "description": "Software",
                "counterparty": "Acme",
                "category": "software",
                "pre_trading": False,
                "tax_treatment": None,
                "source_ref": "receipt:1",
            }],
        }
        decoded = WorkbookCodec.decode_finance(encode_finance_workbook(payload))
        self.assertEqual(payload, decoded)
        self.assertIs(type(decoded["transactions"][0]["amount_pence"]), int)
        self.assertIsNone(decoded["transactions"][0]["tax_treatment"])

    def test_wrong_workbook_kind_is_rejected(self) -> None:
        with self.assertRaises(WorkbookValidationError):
            WorkbookCodec.decode_expense(encode_finance_workbook({
                "schema_version": 1,
                "transactions": [],
            }))

    def test_generated_package_is_deterministic(self) -> None:
        payload = {"schema_version": 1, "transactions": []}
        self.assertEqual(
            encode_finance_workbook(payload),
            encode_finance_workbook(payload),
        )

    def test_update_preserves_comments_widths_and_validations(self) -> None:
        original = encode_expense_workbook({
            "schema_version": 1,
            "expenses": [{"expense_id": "e-1", "amount_pence": 10}],
            "events": [],
            "collisions": [],
            "evidence": [],
        })
        workbook = load_workbook(BytesIO(original))
        sheet = workbook["Expenses"]
        sheet.column_dimensions["A"].width = 41
        sheet["A2"].comment = Comment("operator note", "SEER")
        validation = DataValidation(type="whole")
        validation.add("B2")
        sheet.add_data_validation(validation)
        edited = BytesIO()
        workbook.save(edited)

        updated = WorkbookCodec.update_expense(
            edited.getvalue(),
            {
                "schema_version": 1,
                "expenses": [{"expense_id": "e-1", "amount_pence": 11}],
                "events": [],
                "collisions": [],
                "evidence": [],
            },
        )
        result = load_workbook(BytesIO(updated))["Expenses"]
        self.assertEqual(41, result.column_dimensions["A"].width)
        self.assertEqual("operator note", result["A2"].comment.text)
        self.assertEqual("B2", str(result.data_validations.dataValidation[0].sqref))
        table = next(iter(result.tables.values()))
        self.assertEqual("TableStyleMedium2", table.tableStyleInfo.name)

    def test_expense_update_preserves_approved_receipt_items_sheet(self) -> None:
        original = encode_expense_workbook({
            "schema_version": 1,
            "expenses": [{"expense_id": "e-1", "amount_pence": 10}],
            "events": [],
            "collisions": [],
            "evidence": [],
        })
        workbook = load_workbook(BytesIO(original))
        receipt_items = workbook.create_sheet("Receipt items")
        receipt_items.append(["Receipt ID", "Description"])
        receipt_items.append(["receipt-1", "Conference meal"])
        edited = BytesIO()
        workbook.save(edited)

        decoded = WorkbookCodec.decode_expense(edited.getvalue())
        self.assertEqual(1, len(decoded["expenses"]))

        updated = WorkbookCodec.update_expense(
            edited.getvalue(),
            {
                "schema_version": 1,
                "expenses": [{"expense_id": "e-1", "amount_pence": 11}],
                "events": [],
                "collisions": [],
                "evidence": [],
            },
        )
        result = load_workbook(BytesIO(updated), data_only=False)
        self.assertEqual(
            ["Receipt ID", "Description"],
            [result["Receipt items"]["A1"].value, result["Receipt items"]["B1"].value],
        )
        self.assertEqual("receipt-1", result["Receipt items"]["A2"].value)
        self.assertEqual("Conference meal", result["Receipt items"]["B2"].value)

    def test_expense_update_preserves_custom_xml_package_members_byte_for_byte(self) -> None:
        original = encode_expense_workbook({
            "schema_version": 1,
            "expenses": [{"expense_id": "e-1", "amount_pence": 10}],
            "events": [],
            "collisions": [],
            "evidence": [],
        })
        custom_parts = {
            "customXml/item1.xml": b"<root>one</root>",
            "customXml/item2.xml": b"<root>two</root>",
            "customXml/item3.xml": b"<root>three</root>",
            "customXml/itemProps1.xml": b"<props>one</props>",
            "customXml/itemProps2.xml": b"<props>two</props>",
            "customXml/itemProps3.xml": b"<props>three</props>",
            "customXml/_rels/item1.xml.rels": (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps1.xml"/>'
                b"</Relationships>"
            ),
            "customXml/_rels/item2.xml.rels": (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps2.xml"/>'
                b"</Relationships>"
            ),
            "customXml/_rels/item3.xml.rels": (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps3.xml"/>'
                b"</Relationships>"
            ),
        }
        with ZipFile(BytesIO(original), "r") as source_zip:
            custom_parts["[Content_Types].xml"] = (
                source_zip.read("[Content_Types].xml")
                .replace(
                    b"</Types>",
                    b'<Override PartName="/customXml/item1.xml" ContentType="application/xml"/>'
                    b'<Override PartName="/customXml/item2.xml" ContentType="application/xml"/>'
                    b'<Override PartName="/customXml/item3.xml" ContentType="application/xml"/>'
                    b'<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
                    b'<Override PartName="/customXml/itemProps2.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
                    b'<Override PartName="/customXml/itemProps3.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
                    b"</Types>",
                )
            )
            members = {
                name: source_zip.read(name) for name in source_zip.namelist()
            }
        members.update(custom_parts)
        packaged = BytesIO()
        with ZipFile(packaged, "w", compression=ZIP_DEFLATED) as archive:
            for name, payload in members.items():
                archive.writestr(name, payload)

        updated = WorkbookCodec.update_expense(
            packaged.getvalue(),
            {
                "schema_version": 1,
                "expenses": [{"expense_id": "e-1", "amount_pence": 11}],
                "events": [],
                "collisions": [],
                "evidence": [],
            },
        )
        editable = _editable_package_parts(packaged.getvalue(), "expense")
        with ZipFile(BytesIO(packaged.getvalue()), "r") as before, ZipFile(
            BytesIO(updated), "r"
        ) as after:
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            for name in set(before.namelist()) - editable:
                self.assertEqual(before.read(name), after.read(name), name)
            for name, payload in custom_parts.items():
                self.assertEqual(payload, after.read(name), name)
        self.assertEqual(
            11,
            WorkbookCodec.decode_expense(updated)["expenses"][0]["amount_pence"],
        )

    def test_unapproved_expense_auxiliary_sheet_is_rejected(self) -> None:
        original = encode_expense_workbook({
            "schema_version": 1,
            "expenses": [],
            "events": [],
            "collisions": [],
            "evidence": [],
        })
        workbook = load_workbook(BytesIO(original))
        workbook.create_sheet("Unexpected data")
        edited = BytesIO()
        workbook.save(edited)

        with self.assertRaisesRegex(WorkbookValidationError, "unknown Unexpected data"):
            WorkbookCodec.decode_expense(edited.getvalue())

    def test_blank_marker_accepts_an_ordinary_value_without_marker_edit(self) -> None:
        original = encode_expense_workbook({
            "schema_version": 1,
            "expenses": [{"expense_id": "e-1", "reviewed_amount": None}],
            "events": [],
            "collisions": [],
            "evidence": [],
        })
        workbook = load_workbook(BytesIO(original))
        workbook["Expenses"]["C2"] = 27
        edited = BytesIO()
        workbook.save(edited)
        value = WorkbookCodec.decode_expense(edited.getvalue())
        self.assertEqual(27, value["expenses"][0]["reviewed_amount"])
        self.assertIs(type(value["expenses"][0]["reviewed_amount"]), int)

    def test_formula_edits_fail_closed(self) -> None:
        workbook = load_workbook(BytesIO(encode_finance_workbook({
            "schema_version": 1,
            "transactions": [{"txn_id": "t-1", "amount_pence": 2}],
        })))
        workbook["Transactions"]["C2"] = "=1+1"
        edited = BytesIO()
        workbook.save(edited)
        with self.assertRaises(WorkbookValidationError):
            WorkbookCodec.decode_finance(edited.getvalue())


if __name__ == "__main__":
    unittest.main()
